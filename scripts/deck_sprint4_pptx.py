# -*- coding: utf-8 -*-
"""Monta o .pptx com a transição Transformar entre a tela inteira e os seus destaques.

O efeito que o Igor pediu: ao passar da tela cheia para um destaque, a captura
cresce e desliza até o recorte, em vez de trocar de imagem. Quem assiste entende
sem que ninguém diga que está vendo um pedaço daquela mesma tela.

Como funciona, e por que o slide precisa ser montado de outro jeito por causa
disso:

  · A transição **Transformar** do PowerPoint (Morph, 2016 em diante) anima um
    objeto que existe nos dois slides. Ela casa os objetos pelo nome quando o
    nome começa por `!!`, que é a convenção da própria Microsoft.
  · Um PNG achatado do slide inteiro não serve: para o PowerPoint são duas
    imagens diferentes, e ele faz um esmaecimento. A captura tem de entrar como
    objeto separado, com o mesmo arquivo de origem nos dois slides, mudando só
    a posição, o tamanho e o recorte.
  · Por isso o slide de tela e o de destaque são renderizados **duas vezes**:
    a versão cheia vira o PNG do visualizador, e a versão sem a captura vira o
    fundo do .pptx. Por cima do fundo entra a captura, como retângulo de cantos
    arredondados preenchido pela imagem, cortado por `srcRect`.

Onde não há o que transformar, o slide leva um esmaecimento curto, que é o
comportamento de reserva do próprio XML e o que o LibreOffice entende.

Este módulo é usado pelo `monta_deck_sprint4.py`; ele não roda sozinho.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

# o palco do slide tem 1600 × 900; o slide do PowerPoint, 12192000 × 6858000 EMU
EMU_POR_PX = 12192000 / 1600
LARGURA_EMU, ALTURA_EMU = Emu(12192000), Emu(6858000)
RAIO_PX = 14                      # o mesmo border-radius do cartão no HTML
NOME_MORPH = "!!tela"             # o prefixo `!!` é o que faz o PowerPoint casar os objetos

# duração em milissegundos: a transformação respira, o esmaecimento é curto
DUR_MORPH, DUR_FADE = 900, 350

NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
NS_P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"

_XML_MORPH = (
    '<mc:AlternateContent xmlns:mc="{mc}">'
    '<mc:Choice xmlns:p159="{p159}" Requires="p159">'
    '<p:transition xmlns:p="{p}" xmlns:p14="{p14}" spd="slow" p14:dur="{dur}">'
    '<p159:morph option="byObject"/>'
    "</p:transition></mc:Choice>"
    "<mc:Fallback>"
    '<p:transition xmlns:p="{p}" xmlns:p14="{p14}" spd="med" p14:dur="{fallback}">'
    "<p:fade/></p:transition></mc:Fallback></mc:AlternateContent>"
)
_XML_FADE = (
    '<mc:AlternateContent xmlns:mc="{mc}">'
    '<mc:Choice xmlns:p14="{p14}" Requires="p14">'
    '<p:transition xmlns:p="{p}" spd="med" p14:dur="{dur}">'
    "<p:fade/></p:transition></mc:Choice>"
    "<mc:Fallback>"
    '<p:transition xmlns:p="{p}" spd="med"><p:fade/></p:transition>'
    "</mc:Fallback></mc:AlternateContent>"
)


def _transicao(slide, xml: str, **kw) -> None:
    """Prega a transição no slide.

    A ordem dos filhos de `<p:sld>` é fixa no esquema: `cSld`, `clrMapOvr`,
    `transition`, `timing`. Entrar no lugar errado faz o PowerPoint recusar o
    arquivo, então a inserção é sempre logo depois do `clrMapOvr`.
    """
    from pptx.oxml import parse_xml

    el = parse_xml(xml.format(mc=NS_MC, p=NS_P, p14=NS_P14, p159=NS_P159, **kw))
    sld = slide._element
    anterior = sld.find(qn("p:clrMapOvr"))
    if anterior is None:
        sld.insert(1, el)
    else:
        anterior.addnext(el)


def _arredonda(pic, raio_px: float) -> None:
    """Troca a geometria da imagem de retângulo para retângulo de cantos arredondados.

    O `adj` do PowerPoint é a fração do menor lado, em milésimos de porcento, e é
    assim que o canto acompanha o cartão do HTML em qualquer tamanho.
    """
    from pptx.oxml import parse_xml

    largura = pic.width / EMU_POR_PX
    altura = pic.height / EMU_POR_PX
    adj = int(min(raio_px / min(largura, altura), 0.5) * 100000)
    sp_pr = pic._element.spPr
    antigo = sp_pr.find(qn("a:prstGeom"))
    if antigo is not None:
        sp_pr.remove(antigo)
    geom = parse_xml(
        '<a:prstGeom xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        f'prst="roundRect"><a:avLst><a:gd name="adj" fmla="val {adj}"/></a:avLst></a:prstGeom>'
    )
    # prstGeom vem depois de xfrm e antes do preenchimento
    xfrm = sp_pr.find(qn("a:xfrm"))
    if xfrm is None:
        sp_pr.insert(0, geom)
    else:
        xfrm.addnext(geom)


def _recorta(pic, crop: tuple[float, float, float, float]) -> None:
    """Mostra só o retângulo pedido da imagem, como o `background-position` do HTML."""
    x, y, w, h = crop
    if (x, y, w, h) == (0.0, 0.0, 1.0, 1.0):
        return
    pic.crop_left = max(x, 0.0)
    pic.crop_top = max(y, 0.0)
    pic.crop_right = max(1.0 - x - w, 0.0)
    pic.crop_bottom = max(1.0 - y - h, 0.0)


def monta(pngs: list[Path], notas: dict[int, str], morph: dict[int, dict], destino: Path) -> dict:
    """Escreve o .pptx: um slide por PNG, com a captura solta onde há Transformar.

    `morph[pos]` traz `png_fundo`, `imagem`, `caixa` (x, y, w, h em px do palco),
    `crop` (x, y, w, h em fração da imagem) e `grupo`, que diz a qual tela aquele
    slide pertence. Slides vizinhos do mesmo grupo recebem a transformação.
    """
    prs = Presentation()
    prs.slide_width, prs.slide_height = LARGURA_EMU, ALTURA_EMU
    branco = prs.slide_layouts[6]
    contagem = {"morph": 0, "fade": 0}

    for pos, png in enumerate(pngs, 1):
        info = morph.get(pos)
        s = prs.slides.add_slide(branco)
        fundo = info["png_fundo"] if info else png
        s.shapes.add_picture(str(fundo), 0, 0, width=LARGURA_EMU, height=ALTURA_EMU)

        if info:
            x, y, w, h = info["caixa"]
            pic = s.shapes.add_picture(
                str(info["imagem"]),
                Emu(round(x * EMU_POR_PX)), Emu(round(y * EMU_POR_PX)),
                width=Emu(round(w * EMU_POR_PX)), height=Emu(round(h * EMU_POR_PX)),
            )
            _recorta(pic, info["crop"])
            _arredonda(pic, RAIO_PX)
            pic.name = NOME_MORPH

        # a transformação só vale quando o slide anterior mostra a mesma tela
        anterior = morph.get(pos - 1)
        if info and anterior and anterior["grupo"] == info["grupo"]:
            _transicao(s, _XML_MORPH, dur=DUR_MORPH, fallback=DUR_FADE)
            contagem["morph"] += 1
        else:
            _transicao(s, _XML_FADE, dur=DUR_FADE)
            contagem["fade"] += 1

        s.notes_slide.notes_text_frame.text = notas.get(pos, "")

    prs.save(destino)
    return contagem
